"""One Excel layout for every source: the columns of the card statement."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HEADERS = ["Transaction date", "Post date", "Description", "Amount (VND)", "Type"]
WIDTHS = (16, 12, 60, 18, 8)


def row(tdate, pdate, desc, amount, credit=False):
    """A table row. Credits (CR) get a negative amount."""
    return [tdate, pdate, desc, -amount if credit else amount, "CR" if credit else "DR"]


def write_xlsx(rows, out):
    out = Path(out)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for c in ws["D"][1:]:
        c.number_format = "#,##0.00"
        c.alignment = Alignment(horizontal="right")
    for col, width in zip("ABCDE", WIDTHS):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(out)
    return out
